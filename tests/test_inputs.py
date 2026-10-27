"""The source overlay may change only three approved GWPT cells."""
from pathlib import Path

import pytest
from openpyxl import Workbook

from gweh.manifest import SNAPSHOT_INPUTS, snapshot_files
from gweh.paths import RepositoryPaths
from source import load_overlays, verify_exact_data_changes

ROOT = Path(__file__).resolve().parents[1]


def _workbooks(directory: Path):
    original = directory / "original.xlsx"
    edited = directory / "edited.xlsx"
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["GEM phase ID", "Capacity (MW)", "Status", "other"])
    data.append(["G100001050648", 461.6, "pre-construction", "untouched"])
    data.append(["G100000927160", 336.6, "construction", "untouched"])
    data.append(["G100001050672", 317.0, "pre-construction", "untouched"])
    data.append(["CONTROL", 20.0, "operating", "untouched"])
    original_sheet = workbook.create_sheet("Other")
    original_sheet.append(["keep"])
    workbook.save(original)
    data["B2"] = 576
    data["B3"] = 306
    data["C4"] = "cancelled"
    workbook.save(edited)
    return original, edited


def test_approved_three_cell_source_overlay(tmp_path):
    original, edited = _workbooks(tmp_path)
    overlay = load_overlays(ROOT / "config/gwpt_source_overrides.csv")
    changed, _ = verify_exact_data_changes(original, edited, overlay)
    assert {(row["phase_id"], row["field"]) for row in changed} == {
        (row["phase_id"], row["field"]) for row in overlay
    }


def test_fourth_source_change_is_rejected(tmp_path):
    original, edited = _workbooks(tmp_path)
    from openpyxl import load_workbook
    workbook = load_workbook(edited)
    workbook["Data"]["B5"] = 21.0
    workbook.save(edited)
    overlay = load_overlays(ROOT / "config/gwpt_source_overrides.csv")
    with pytest.raises(RuntimeError, match="Unexpected workbook value changed"):
        verify_exact_data_changes(original, edited, overlay)


def test_snapshot_contains_only_numeric_inputs(tmp_path):
    files = snapshot_files(RepositoryPaths.discover(ROOT), tmp_path)
    assert set(files) == set(SNAPSHOT_INPUTS)
    assert not (tmp_path / "provenance").exists()
    assert not (tmp_path / "evidence/strategic_s/source_database.csv").exists()


def test_source_override_contract_rejects_historical_columns(tmp_path):
    import pandas as pd
    ref = ROOT / "config/gwpt_source_overrides.csv"
    rows = pd.read_csv(ref, dtype=str)
    assert len(rows) == 3
    assert "as_of" not in rows.columns
    bad = tmp_path / "unwanted.csv"
    rows.assign(as_of="2026-02").to_csv(bad, index=False)
    with pytest.raises(ValueError, match="unexpected fields"):
        load_overlays(bad)
