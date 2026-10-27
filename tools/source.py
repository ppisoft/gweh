#!/usr/bin/env python3
"""Prepare a separate GWPT input with exactly three documented corrections."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from itertools import zip_longest
from pathlib import Path

from openpyxl import load_workbook

ORIGINAL_GWPT_SHA256 = "5d828a1b5d9a2b862aa009c23639e05fb2b3e0ad14baf0c8c9b83cc115a43210"
SOURCE_NAME = "Global-Wind-Power-Tracker-February-2026.xlsx"
SCENARIO_DEFAULT = "gweh-source"
APPROVED_CELLS = {
    ("G100001050648", "Capacity (MW)"),
    ("G100000927160", "Capacity (MW)"),
    ("G100001050672", "Status"),
}


def sha(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(8 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_overlays(path: Path) -> list[dict[str, str]]:
    with Path(path).open(newline="", encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        expected = ["phase_id", "field", "expected_original", "effective_value",
                    "source_alias", "source_date", "ontology", "reason"]
        if reader.fieldnames != expected:
            raise ValueError("GWPT source overrides have unexpected fields")
        rows = list(reader)
    keys = [(row["phase_id"], row["field"]) for row in rows]
    if len(rows) != len(APPROVED_CELLS) or set(keys) != APPROVED_CELLS:
        raise ValueError("GWPT correction file must declare exactly the three approved cells")
    return rows


def verify_exact_data_changes(old_path: Path, new_path: Path, overlays):
    """Compare every value in every worksheet; reject any extra changed cell."""
    allowed = {
        (item["phase_id"], item["field"]): (item["expected_original"], item["effective_value"])
        for item in overlays
    }
    original = load_workbook(old_path, read_only=True, data_only=True)
    modified = load_workbook(new_path, read_only=True, data_only=True)
    try:
        if original.sheetnames != modified.sheetnames:
            raise RuntimeError("Sheet list changed")
        changes = []
        sheet_counts = {}
        sentinel = object()
        for name in original.sheetnames:
            old_rows = original[name].iter_rows(values_only=True)
            new_rows = modified[name].iter_rows(values_only=True)
            original_header = next(old_rows)
            modified_header = next(new_rows)
            if original_header != modified_header:
                raise RuntimeError(f"Header changed: {name}")
            column_by_name = {str(header): idx for idx, header in enumerate(original_header)}
            phase_index = column_by_name.get("GEM phase ID")
            count = 1
            for original_row, modified_row in zip_longest(old_rows, new_rows, fillvalue=sentinel):
                count += 1
                if original_row is sentinel or modified_row is sentinel:
                    raise RuntimeError(f"Row count changed: {name}")
                if len(original_row) != len(modified_row):
                    raise RuntimeError(f"Column count changed: {name}/{count}")
                phase = (
                    str(original_row[phase_index])
                    if phase_index is not None and name == "Data" else ""
                )
                for idx, (before, after) in enumerate(zip(original_row, modified_row)):
                    if before == after:
                        continue
                    key = phase, str(original_header[idx])
                    if name != "Data" or key not in allowed:
                        raise RuntimeError(
                            f"Unexpected workbook value changed: {name}/{count}/{key}: "
                            f"{before!r} -> {after!r}"
                        )
                    expected, replacement = allowed[key]
                    same_replacement = str(after) == str(replacement)
                    if isinstance(after, (float, int)):
                        same_replacement |= float(after) == float(replacement)
                    if str(before) != str(expected) or not same_replacement:
                        raise RuntimeError(f"Incorrect GWPT correction at {key}")
                    changes.append({
                        "sheet": name, "row": count, "phase_id": phase,
                        "field": key[1], "from": str(before), "to": str(after),
                    })
            sheet_counts[name] = count
        if len(changes) != len(allowed) or {
            (item["phase_id"], item["field"]) for item in changes
        } != set(allowed):
            raise RuntimeError("Exactly three approved GWPT changes required")
        return changes, sheet_counts
    finally:
        original.close()
        modified.close()


def _scenario_path(parent: Path, scenario_id: str) -> Path:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", scenario_id) or ".." in scenario_id:
        raise ValueError("Use a simple scenario identifier without path separators")
    return parent / "science_scenarios" / scenario_id


def prepare(args):
    repo = Path(args.repo).resolve()
    root = Path(args.data_root).resolve()
    original = (
        Path(args.original_workbook).resolve()
        if args.original_workbook else root / "raw/GWPT" / SOURCE_NAME
    )
    overlays = load_overlays(repo / "config/gwpt_source_overrides.csv")
    if not original.is_file():
        raise FileNotFoundError(original)
    if sha(original) != ORIGINAL_GWPT_SHA256:
        raise RuntimeError("Original GWPT source SHA256 does not match the February 2026 workbook")
    scenario = _scenario_path(root, args.scenario_id)
    if scenario.exists():
        raise FileExistsError(f"Scenario already exists: {scenario}")
    if args.dry_run:
        print(f"[ready] original GWPT and three-cell correction list: {scenario}")
        return

    temporary = scenario.parent / f".{args.scenario_id}.inprogress"
    if temporary.exists():
        raise FileExistsError(f"Inspect the existing incomplete staging directory: {temporary}")
    temporary.mkdir(parents=True)
    try:
        data = temporary / "data"
        edited = data / "raw/GWPT" / SOURCE_NAME
        edited.parent.mkdir(parents=True, exist_ok=True)
        workbook = load_workbook(original, read_only=False, data_only=False)
        worksheet = workbook["Data"]
        columns = {str(cell.value): cell.column for cell in worksheet[1]}
        for field in ("GEM phase ID", "Capacity (MW)", "Status"):
            if field not in columns:
                raise RuntimeError(f"Original workbook lacks {field}")

        target_ids = {item["phase_id"] for item in overlays}
        indices = {}
        for row_number in range(2, worksheet.max_row + 1):
            phase = str(worksheet.cell(row_number, columns["GEM phase ID"]).value)
            if phase not in target_ids:
                continue
            if phase in indices:
                raise RuntimeError(f"Duplicate target phase ID {phase}")
            indices[phase] = row_number
        if set(indices) != target_ids:
            raise RuntimeError("Some approved phase IDs are absent from the original GWPT")

        for item in overlays:
            cell = worksheet.cell(indices[item["phase_id"]], columns[item["field"]])
            if str(cell.value) != item["expected_original"]:
                raise RuntimeError(f"GWPT correction precondition failed: {item['phase_id']}")
            cell.value = (
                float(item["effective_value"])
                if item["field"] == "Capacity (MW)" else item["effective_value"]
            )
        workbook.save(edited)
        workbook.close()
        applied, sheet_counts = verify_exact_data_changes(original, edited, overlays)

        if not args.source_only:
            for name in ("EEZ", "NaturalEarth", "GWA"):
                source = root / "raw" / name
                if not source.is_dir():
                    raise FileNotFoundError(source)
                (data / "raw" / name).symlink_to(source, target_is_directory=True)
        (data / "runs").mkdir()
        audit = {
            "status": "PASS_ONLY_THREE_EXPECTED_GWPT_VALUE_CHANGES",
            "scenario_id": args.scenario_id,
            "source_sha256": ORIGINAL_GWPT_SHA256,
            "source_workbook": str(original),
            "new_source_sha256": sha(edited),
            "updated_cells": applied,
            "sheets_checked": sheet_counts,
            "scenario_workbook_relative_path": f"raw/GWPT/{SOURCE_NAME}",
        }
        (temporary / "SOURCE_OVERLAY_AUDIT.json").write_text(
            json.dumps(audit, indent=2) + "\n", encoding="utf-8"
        )
        temporary.rename(scenario)
        print(f"[done] prepared three-cell GWPT scenario: {scenario}")
    except BaseException:
        print(f"[failed] inspect staging directory: {temporary}", file=sys.stderr)
        raise


def verify_scenario(repo: Path, parent: Path, scenario_id: str, *, source_only=False):
    parent = parent.resolve()
    scenario = _scenario_path(parent, scenario_id)
    audit = json.loads((scenario / "SOURCE_OVERLAY_AUDIT.json").read_text())
    if audit.get("status") != "PASS_ONLY_THREE_EXPECTED_GWPT_VALUE_CHANGES":
        raise RuntimeError("Scenario correction audit is incomplete")
    original = Path(audit.get("source_workbook", str(parent / "raw/GWPT" / SOURCE_NAME)))
    modified = scenario / "data/raw/GWPT" / SOURCE_NAME
    if sha(original) != ORIGINAL_GWPT_SHA256 or sha(modified) != audit["new_source_sha256"]:
        raise RuntimeError("Original or prepared GWPT workbook changed")
    expected = {
        (row["phase_id"], row["field"], row["expected_original"], row["effective_value"])
        for row in load_overlays(repo / "config/gwpt_source_overrides.csv")
    }
    observed = {
        (row["phase_id"], row["field"], row["from"], row["to"])
        for row in audit["updated_cells"]
    }
    if observed != expected:
        raise RuntimeError("GWPT correction audit differs from the approved three-cell list")
    if not source_only:
        for name in ("EEZ", "NaturalEarth", "GWA"):
            link = scenario / "data/raw" / name
            if not link.is_symlink() or link.resolve() != (parent / "raw" / name).resolve():
                raise RuntimeError(f"GIS source link changed: {name}")
    print(f"[PASS] prepared GWPT scenario: {scenario}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, required=True)
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--original-workbook")
    parser.add_argument("--scenario-id", default=SCENARIO_DEFAULT)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--source-only", action="store_true", help="Workbook-only unit test")
    parser.add_argument("--verify", action="store_true", help="Verify an existing scenario")
    args = parser.parse_args()
    if args.verify:
        verify_scenario(args.repo, args.data_root, args.scenario_id, source_only=args.source_only)
    else:
        prepare(args)


if __name__ == "__main__":
    main()
